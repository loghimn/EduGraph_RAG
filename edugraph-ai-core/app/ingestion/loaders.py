import os
import pymupdf as fitz  # PyMuPDF
import docx
from pptx import Presentation
import re
from typing import List, Tuple
from app.config import settings
from app.core.errors import InvalidPathError, ScannedPDFError, UnsupportedFormatError

def get_safe_path(doc_id: str, filename: str) -> str:
    base_dir = os.path.abspath(settings.RAW_DATA_DIR)
    target_path = os.path.abspath(os.path.join(base_dir, doc_id, filename))
    if not target_path.startswith(base_dir):
        raise InvalidPathError(f"Path traversal detected: {target_path}")
    return target_path

def load_pdf(path: str) -> List[Tuple[int, str]]:
    """Trả về danh sách (page_number, text). Ném lỗi nếu là PDF scan."""
    doc = fitz.open(path)
    num_pages = len(doc)
    if num_pages == 0:
        raise ScannedPDFError("Tài liệu PDF trống.")
        
    pages = []
    total_chars = 0
    for i in range(num_pages):
        text = doc[i].get_text("text") or ""
        total_chars += len(text.strip())
        pages.append((i + 1, text))
        
    avg_chars = total_chars / num_pages
    if avg_chars < settings.MIN_CHARS_PER_PAGE:
        raise ScannedPDFError(f"Tài liệu có vẻ là bản scan (Trung bình {avg_chars:.1f} ký tự/trang).")
        
    return pages

def load_docx(path: str) -> List[Tuple[int, str]]:
    """Đọc DOCX, lấy cả text và text trong bảng."""
    doc = docx.Document(path)
    text_blocks = []
    
    # Đoạn văn
    for para in doc.paragraphs:
        if para.text.strip():
            text_blocks.append(para.text)
            
    # Bảng
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                text_blocks.append(" | ".join(row_text))
                
    # Gộp lại và vờ như nó ở trang 1
    return [(1, "\n".join(text_blocks))]

def load_pptx(path: str) -> List[Tuple[int, str]]:
    """Đọc PPTX — mỗi slide là một "trang", lấy text từ shapes (title, body, bảng)."""
    prs = Presentation(path)
    num_slides = len(prs.slides)
    if num_slides == 0:
        raise ScannedPDFError("Tài liệu PPTX trống.")

    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        blocks = []
        for shape in slide.shapes:
            # Bảng (GraphicFrame không có text_frame)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        blocks.append(" | ".join(row_text))
                continue
            # Shape có text frame (title, placeholder, textbox...)
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text
                if text.strip():
                    blocks.append(text)
        pages.append((i, "\n".join(blocks)))

    return pages

def load_text(path: str) -> List[Tuple[int, str]]:
    """Đọc TXT / MD"""
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()
    return [(1, text)]

def load_document(doc_id: str, filename: str) -> List[Tuple[int, str]]:
    path = get_safe_path(doc_id, filename)
    if not os.path.exists(path):
        raise FileNotFoundError(f"Không tìm thấy file: {path}")
        
    ext = filename.lower().split('.')[-1]
    if ext == 'pdf':
        return load_pdf(path)
    elif ext == 'docx':
        return load_docx(path)
    elif ext == 'pptx':
        return load_pptx(path)
    elif ext == 'doc':
        raise UnsupportedFormatError("Định dạng .doc cũ không được hỗ trợ. Vui lòng chuyển sang .docx.")
    elif ext in ['md', 'txt']:
        return load_text(path)
    else:
        raise UnsupportedFormatError(f"Không hỗ trợ định dạng .{ext}")
