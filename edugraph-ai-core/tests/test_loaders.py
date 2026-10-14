import pytest
from pptx import Presentation

from app.ingestion.loaders import load_document, load_pptx
from app.core.errors import ScannedPDFError


def _make_pptx(path, slides):
    """Tạo file .pptx thật với các slide (title, body)."""
    prs = Presentation()
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(str(path))


def test_load_pptx_extracts_each_slide(tmp_path):
    path = tmp_path / "slides.pptx"
    _make_pptx(path, [("Giới thiệu", "Slide nội dung 1"), ("Kết luận", "Slide nội dung 2")])

    pages = load_pptx(str(path))

    assert [p for p, _ in pages] == [1, 2]
    assert "Giới thiệu" in pages[0][1]
    assert "Slide nội dung 1" in pages[0][1]
    assert "Kết luận" in pages[1][1]
    assert "Slide nội dung 2" in pages[1][1]


def test_load_document_dispatch_pptx(tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "RAW_DATA_DIR", str(tmp_path))

    doc_dir = tmp_path / "doc-pptx-1"
    doc_dir.mkdir()
    _make_pptx(doc_dir / "bai2.pptx", [("Chương 1", "Nội dung chương 1")])

    pages = load_document("doc-pptx-1", "bai2.pptx")

    assert len(pages) == 1
    assert pages[0][0] == 1
    assert "Chương 1" in pages[0][1]


def test_load_pptx_empty_raises(tmp_path):
    path = tmp_path / "empty.pptx"
    Presentation().save(str(path))

    with pytest.raises(ScannedPDFError):
        load_pptx(str(path))
